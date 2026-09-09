from __future__ import annotations

import argparse
import importlib
import json
import os
import sys
import traceback
from pathlib import Path


PROFILE_MODULES = {
    "base": [
        "numpy", "pandas", "scipy", "matplotlib", "sympy", "jupyterlab",
        "openpyxl", "xlrd", "PIL",
    ],
    "optimization": ["cvxpy", "pulp", "ortools", "networkx"],
    "ml": ["sklearn", "statsmodels", "xgboost"],
    "paper": ["docx", "fitz", "pypdf"],
}

PROFILE_CLOSURE = {
    "base": ["base"],
    "optimization": ["base", "optimization"],
    "ml": ["base", "ml"],
    "paper": ["base", "paper"],
    "full": ["base", "optimization", "ml", "paper"],
}


def result(name: str, function) -> dict[str, str]:
    try:
        detail = function()
        return {"name": name, "status": "PASS", "detail": str(detail)}
    except Exception as exc:  # the report must retain all failures
        return {
            "name": name,
            "status": "FAIL",
            "detail": f"{type(exc).__name__}: {exc}",
            "traceback": traceback.format_exc(limit=4),
        }


def check_imports(profile: str) -> str:
    modules: list[str] = []
    for group in PROFILE_CLOSURE[profile]:
        modules.extend(PROFILE_MODULES[group])
    for name in modules:
        importlib.import_module(name)
    return f"成功导入 {len(modules)} 个模块：{', '.join(modules)}"


def check_scientific_computing() -> str:
    import numpy as np
    from scipy.optimize import linprog

    optimum = linprog(c=[-3, -2], A_ub=[[1, 1], [2, 1]], b_ub=[4, 6], bounds=(0, None))
    if not optimum.success or not np.isclose(-optimum.fun, 10.0):
        raise RuntimeError(f"线性规划结果异常：{optimum.message}")
    return f"SciPy 线性规划最优值={-optimum.fun:.6g}"


def check_excel(output_dir: Path) -> str:
    import pandas as pd

    path = output_dir / "modelready-check.xlsx"
    expected = pd.DataFrame({"方案": ["甲", "乙"], "得分": [88.5, 91.0]})
    expected.to_excel(path, index=False)
    actual = pd.read_excel(path)
    if not actual.equals(expected):
        raise RuntimeError("Excel 写入后读回的数据不一致")
    return f"Excel 往返读写通过：{path.name}"


def check_plot(output_dir: Path) -> str:
    os.environ.setdefault("MPLBACKEND", "Agg")
    import matplotlib.pyplot as plt
    import numpy as np
    from matplotlib import font_manager, rcParams

    preferred_fonts = [
        "Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "Source Han Sans SC",
        "WenQuanYi Zen Hei",
    ]
    installed_fonts = {font.name for font in font_manager.fontManager.ttflist}
    selected_font = next((name for name in preferred_fonts if name in installed_fonts), None)
    if selected_font is None:
        raise RuntimeError("未找到受支持的中文字体（微软雅黑、黑体或 Noto/Source Han CJK）")
    rcParams["font.sans-serif"] = [selected_font]
    rcParams["axes.unicode_minus"] = False

    path = output_dir / "modelready-check.png"
    x = np.linspace(0, 2 * np.pi, 100)
    fig, ax = plt.subplots(figsize=(5, 3))
    ax.plot(x, np.sin(x), label="sin(x)")
    ax.set_title("ModelReady 中文绘图验证")
    ax.legend()
    fig.tight_layout()
    fig.savefig(path, dpi=160)
    plt.close(fig)
    if path.stat().st_size < 1_000:
        raise RuntimeError("绘图文件异常小")
    return f"生成绘图：{path.name} ({path.stat().st_size} bytes)，字体={selected_font}"


def check_machine_learning() -> str:
    import numpy as np
    from sklearn.linear_model import LinearRegression

    x = np.arange(10, dtype=float).reshape(-1, 1)
    y = 2.5 * x.ravel() + 4
    model = LinearRegression().fit(x, y)
    if not np.isclose(model.coef_[0], 2.5) or model.score(x, y) < 0.999999:
        raise RuntimeError("线性回归验证未达到预期")
    return f"线性回归 R²={model.score(x, y):.6f}"


def check_optimization() -> str:
    import cvxpy as cp

    x = cp.Variable(2, nonneg=True)
    problem = cp.Problem(cp.Maximize(3 * x[0] + 2 * x[1]), [cp.sum(x) <= 4, 2 * x[0] + x[1] <= 6])
    value = problem.solve()
    if value is None or abs(value - 10) > 1e-4:
        raise RuntimeError(f"CVXPY 最优值异常：{value}")
    return f"CVXPY 最优值={value:.6g}，求解器={problem.solver_stats.solver_name}"


def check_document_packages(output_dir: Path) -> str:
    from docx import Document
    from pypdf import PdfReader  # noqa: F401 - import validates package availability

    path = output_dir / "modelready-check.docx"
    document = Document()
    document.add_heading("ModelReady 文档验证", 0)
    document.add_paragraph("数学建模环境已能够生成 Word 文档。")
    document.save(path)
    if path.stat().st_size < 5_000:
        raise RuntimeError("Word 文档异常小")
    return f"生成 Word 文档：{path.name}"


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--profile", choices=PROFILE_CLOSURE, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    results = [
        result("依赖导入", lambda: check_imports(args.profile)),
        result("科学计算与线性规划", check_scientific_computing),
        result("Excel 往返读写", lambda: check_excel(args.output_dir)),
        result("中文绘图", lambda: check_plot(args.output_dir)),
    ]
    closure = PROFILE_CLOSURE[args.profile]
    if "ml" in closure:
        results.append(result("机器学习", check_machine_learning))
    if "optimization" in closure:
        results.append(result("凸优化", check_optimization))
    if "paper" in closure:
        results.append(result("Word/PDF 包", lambda: check_document_packages(args.output_dir)))

    payload = {
        "schemaVersion": 1,
        "profile": args.profile,
        "python": sys.version,
        "executable": sys.executable,
        "results": results,
    }
    output = args.output_dir / f"verify-{args.profile}.json"
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(output)
    return 0 if all(item["status"] == "PASS" for item in results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
